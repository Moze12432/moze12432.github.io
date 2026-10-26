#!/usr/bin/env python3
"""MozeAI reproducible benchmark harness (IJASC revision, build ijasc-r1).

  python3 -u mozeai_bench.py local  --app emotion-ai-demo/app.py --out bench_results
  python3 -u mozeai_bench.py online --app emotion-ai-demo/app.py --out bench_results   # needs GROQ_API_KEY
  python3 -u mozeai_bench.py online --app emotion-ai-demo/app.py --out bench_results --dry-run   # offline self-test

The harness loads the functions of app.py verbatim (all top-level def/class blocks,
constants and session-state initialisers) with a no-op Streamlit stub, so every
measurement exercises the application's own code paths.
"""
import argparse, ast, contextlib, csv, gc, hashlib, heapq, io, json, os, platform, random, re
import statistics, subprocess, sys, time, types
from collections import defaultdict
from importlib import metadata

ALPACA_URL = ("https://raw.githubusercontent.com/tatsu-lab/stanford_alpaca/"
              "761dc5bfbdeeffa89b8bff5d038781a4055f796a/alpaca_data.json")
ALPACA_SHA256 = "2eddafc6b977608d778aaab8dfc7e50e547b3af9826dfb9e909d9fc362e4a419"
SEED = 2026
UNAVAILABLE = "AI service temporarily unavailable. Please try again."
PKGS = ["streamlit", "groq", "requests", "beautifulsoup4", "PyPDF2", "python-docx", "python-pptx",
        "openpyxl", "reportlab", "pytz", "numpy", "matplotlib"]


def log(*a):
    print(time.strftime("%H:%M:%S"), *a, flush=True)


# ---------------------------------------------------------------- environment / data
def environment(app_path):
    env = {"utc_time": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
           "platform": platform.platform(), "machine": platform.machine(),
           "python": platform.python_version(), "logical_cpus": os.cpu_count()}
    cpu, ram = platform.processor(), None
    try:
        with open("/proc/cpuinfo") as f:
            cpu = next(l.split(":", 1)[1].strip() for l in f if l.startswith("model name"))
        with open("/proc/meminfo") as f:
            ram = round(int(next(l.split()[1] for l in f if l.startswith("MemTotal"))) / 1048576, 1)
    except Exception:
        try:
            cpu = subprocess.run(["sysctl", "-n", "machdep.cpu.brand_string"], capture_output=True, text=True).stdout.strip() or cpu
            ram = round(int(subprocess.run(["sysctl", "-n", "hw.memsize"], capture_output=True, text=True).stdout) / 2**30, 1)
        except Exception:
            pass
    env.update(cpu=cpu, ram_gb=ram)
    vers = {}
    for p in PKGS:
        try:
            vers[p] = metadata.version(p)
        except Exception:
            vers[p] = None
    env["packages"] = vers
    src = open(app_path, "rb").read().replace(b"\r\n", b"\n")
    env["app_sha256"] = hashlib.sha256(src).hexdigest()
    try:
        d = os.path.dirname(os.path.abspath(app_path))
        env["git_commit"] = subprocess.run(["git", "-C", d, "rev-parse", "HEAD"], capture_output=True, text=True).stdout.strip()
        env["git_describe"] = subprocess.run(["git", "-C", d, "describe", "--tags", "--always", "--dirty"], capture_output=True, text=True).stdout.strip()
    except Exception:
        pass
    return env


def load_alpaca(out):
    path = os.path.join(out, "data", "alpaca_data.json")
    os.makedirs(os.path.dirname(path), exist_ok=True)
    if not os.path.exists(path):
        import requests
        log("downloading Alpaca dataset (pinned commit 761dc5b, 22 MB) ...")
        r = requests.get(ALPACA_URL, timeout=120)
        r.raise_for_status()
        open(path, "wb").write(r.content)
    digest = hashlib.sha256(open(path, "rb").read()).hexdigest()
    if digest != ALPACA_SHA256:
        sys.exit(f"Alpaca checksum mismatch: {digest}")
    data = json.load(open(path, encoding="utf-8"))
    seen, pairs = set(), []
    for x in data:  # add_memory() ignores texts shorter than 50 characters
        if len(x["output"]) >= 50 and x["output"] not in seen:
            seen.add(x["output"])
            pairs.append(x)
    return data, pairs


# ---------------------------------------------------------------- Streamlit stub + loader
class SessionState(dict):
    def __getattr__(self, k):
        try:
            return self[k]
        except KeyError:
            raise AttributeError(k)

    def __setattr__(self, k, v):
        self[k] = v

    def __delattr__(self, k):
        del self[k]


class StStub:
    def __init__(self):
        self.session_state = SessionState()
        self.secrets = {}

    @contextlib.contextmanager
    def spinner(self, *a, **k):
        yield

    def __getattr__(self, name):
        return lambda *a, **k: None


def load_app(app_path, client=None):
    src = open(app_path, encoding="utf-8").read()
    tree = ast.parse(src)
    imports, consts, defs, session = [], [], [], []
    for n in tree.body:
        s = ast.get_source_segment(src, n)
        if isinstance(n, (ast.FunctionDef, ast.ClassDef)):
            defs.append(s)
        elif isinstance(n, (ast.Import, ast.ImportFrom)):
            mods = [a.name for a in n.names] if isinstance(n, ast.Import) else [n.module or ""]
            if not any(m.split(".")[0] in ("streamlit", "groq") for m in mods):
                imports.append(s)
        elif isinstance(n, ast.Assign) and any(isinstance(t, ast.Name) and t.id in ("TEMPERATURE", "MAX_TOKENS", "SYSTEM_PROMPT") for t in n.targets):
            consts.append(s)
        elif isinstance(n, ast.If) and "not in st.session_state" in ast.unparse(n.test):
            session.append(s)
        elif isinstance(n, ast.Assign) and any(isinstance(t, ast.Name) and t.id == "memory" for t in n.targets):
            session.append(s)
    ns = {"__name__": "mozeai_app", "st": StStub(), "client": client}
    exec(compile("\n".join(imports + consts), app_path, "exec"), ns)
    exec(compile("\n\n".join(defs), app_path, "exec"), ns)
    session_code = compile("\n".join(session), app_path, "exec")

    def reset_session():
        ns["st"].session_state.clear()
        exec(session_code, ns)

    reset_session()
    chain = []
    fn = next(n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name == "llm_with_fallback")
    for node in ast.walk(fn):
        if isinstance(node, ast.Assign) and any(isinstance(t, ast.Name) and t.id == "models_to_try" for t in node.targets):
            chain = [ast.literal_eval(e) for e in node.value.elts]
    ns["_SRC"] = src
    return ns, reset_session, chain


class Upload(io.BytesIO):
    """Mimics Streamlit's UploadedFile (BytesIO with .name and .type)."""
    def __init__(self, data, name, mime):
        super().__init__(data)
        self.name, self.type = name, mime


MIME = {"pdf": "application/pdf", "txt": "text/plain", "csv": "text/csv", "json": "application/json",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}


# ---------------------------------------------------------------- deterministic test files
def make_test_files(out, pairs):
    from docx import Document
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    d = os.path.join(out, "data", "files")
    os.makedirs(d, exist_ok=True)
    text = "\n\n".join(p["output"] for p in pairs)
    files = []

    def add(name, data, kind, scale):
        p = os.path.join(d, name)
        open(p, "wb").write(data)
        files.append({"name": name, "path": p, "type": kind, "scale": scale, "bytes": len(data)})

    for kb in (10, 100, 1024):
        add(f"text_{kb}KB.txt", text.encode("utf-8")[:kb * 1024].decode("utf-8", "ignore").encode("utf-8"), "txt", f"{kb} KB")
    for n in (100, 1000, 10000):
        s = io.StringIO()
        w = csv.writer(s)
        w.writerow(["id", "instruction", "output_words", "output_chars"])
        for i, p in enumerate(pairs[:n]):
            w.writerow([i + 1, p["instruction"], len(p["output"].split()), len(p["output"])])
        add(f"table_{n}rows.csv", s.getvalue().encode("utf-8"), "csv", f"{n} rows")
    for n in (100, 1000, 10000):
        recs = [{"id": i + 1, "instruction": p["instruction"], "input": p["input"], "output": p["output"]} for i, p in enumerate(pairs[:n])]
        add(f"records_{n}.json", json.dumps(recs).encode("utf-8"), "json", f"{n} records")
    for n in (5, 50, 500):
        doc = Document()
        import datetime as _dt
        doc.core_properties.created = doc.core_properties.modified = _dt.datetime(2026, 1, 1)
        for p in pairs[:n]:
            doc.add_paragraph(p["output"])
        t = doc.add_table(rows=10, cols=4)
        for r in range(10):
            for c in range(4):
                t.cell(r, c).text = f"R{r + 1}C{c + 1}"
        b = io.BytesIO()
        doc.save(b)
        add(f"document_{n}para.docx", b.getvalue(), "docx", f"{n} paragraphs")
    words = text.split()
    for pages in (1, 10, 50):
        b = io.BytesIO()
        c = canvas.Canvas(b, pagesize=A4, invariant=1)
        k = 0
        for _ in range(pages):
            y = 800
            for _ in range(45):
                line = []
                while len(" ".join(line)) < 85 and k < len(words):
                    line.append(words[k])
                    k += 1
                c.drawString(50, y, " ".join(line))
                y -= 16
            c.showPage()
        c.save()
        add(f"report_{pages}pages.pdf", b.getvalue(), "pdf", f"{pages} pages")
    return files


def pctl(v, q):
    v = sorted(v)
    if not v:
        return None
    i = (len(v) - 1) * q
    lo, hi = int(i), min(int(i) + 1, len(v) - 1)
    return v[lo] + (v[hi] - v[lo]) * (i - lo)


def stats_ms(ns_list):
    v = [x / 1e6 for x in ns_list]
    return {"n": len(v), "median_ms": statistics.median(v), "q1_ms": pctl(v, .25), "q3_ms": pctl(v, .75),
            "p95_ms": pctl(v, .95), "mean_ms": statistics.fmean(v)}


# ---------------------------------------------------------------- LOCAL benchmarks
def timeit_ns(fn, warmup=3, reps=30):
    for _ in range(warmup):
        fn()
    out = []
    gc.collect()
    gc.disable()
    try:
        for _ in range(reps):
            t = time.perf_counter_ns()
            fn()
            out.append(time.perf_counter_ns() - t)
    finally:
        gc.enable()
    return out


def bench_parsing(ns, files, reps):
    rows = []
    for f in files:
        data = open(f["path"], "rb").read()
        res = {}

        def run():
            res["text"] = ns["process_uploaded_file"](Upload(data, f["name"], MIME[f["type"]]))
        t = timeit_ns(run, 3, reps)
        st = stats_ms(t)
        if res["text"].startswith(("Error", "Unsupported")):
            sys.exit(f"parser failed on {f['name']}: {res['text'][:120]}")
        rows.append({**{k: f[k] for k in ("name", "type", "scale", "bytes")}, **st, "extracted_chars": len(res["text"])})
        log(f"parse {f['name']:<26} {f['bytes'] / 1024:9.1f} KB  median {st['median_ms']:8.2f} ms  IQR {st['q1_ms']:.2f}-{st['q3_ms']:.2f}")
    return rows


PPT_CONTENT = """Introduction to Smart Campus Assistants
- Students need fast answers about courses, schedules and services
- Assistants combine search, documents and conversation in one place
- Open-source tools make such assistants affordable for universities

Key Features
- Real-time web search returns current information with sources
- Uploaded PDF, Word, CSV and JSON files can be summarised or compared
- Presentations, Word reports and Excel sheets are generated on request

Benefits/Importance
- Saves time on routine information requests
- Lowers the barrier to producing structured documents
- Runs on free cloud tiers without dedicated GPUs

Conclusion
- Integration matters more than any single capability
- Reliability requires fallback across several hosted models
- Memory and retrieval quality remain open research problems"""


def bench_assembly(ns, reset, pairs, reps):
    rows = []
    word_text = "\n\n".join(p["output"] for p in pairs[100:110])
    cases = [("PowerPoint (.pptx)", "Smart Campus Assistants", lambda: ns["create_ppt_from_content"]("Smart Campus Assistants", PPT_CONTENT)),
             ("Word (.docx)", "Smart Campus Assistants", lambda: ns["create_word_from_content"]("Smart Campus Assistants", word_text))]
    for label, _, fn in cases:
        res = {}
        t = timeit_ns(lambda: res.__setitem__("b", fn()), 3, reps)
        size = len(res["b"].getvalue())
        st = stats_ms(t)
        rows.append({"artifact": label, "input": f"{len(PPT_CONTENT.split()) if 'Power' in label else len(word_text.split())} words", "bytes": size, **st})
        log(f"build {label:<22} median {st['median_ms']:7.2f} ms  ({size / 1024:.1f} KB)")
    for prompt in ("generate an excel about phone sales", "generate an excel about student data", "generate an excel about project tasks"):
        res = {}

        def run():
            reset()
            res["r"] = ns["run_agent"](prompt)
            res["b"] = ns["st"].session_state.get("excel_data")
        t = timeit_ns(run, 3, reps)
        if res["b"] is None:
            sys.exit(f"Excel generation failed for: {prompt}")
        st = stats_ms(t)
        size = len(res["b"].getvalue())
        rows.append({"artifact": f"Excel (.xlsx): {prompt.split('about ')[1]}", "input": "template", "bytes": size, **st})
        log(f"build Excel ({prompt.split('about ')[1]:<13}) median {st['median_ms']:7.2f} ms  ({size / 1024:.1f} KB)")
    return rows


def memory_class(ns, cap):
    src = ns["_SRC"]
    tree = ast.parse(src)
    cls = next(n for n in tree.body if isinstance(n, ast.ClassDef) and n.name == "LightweightMemory")
    s = ast.get_source_segment(src, cls)
    s2, k1 = re.subn(r"len\(self\.memories\) > 50", "len(self.memories) > _CAP", s)
    s2, k2 = re.subn(r"self\.memories\[-50:\]", "self.memories[-_CAP:]", s2)
    if (k1, k2) != (1, 1):
        sys.exit("could not locate the 50-entry window in LightweightMemory")
    g = dict(ns)
    g["_CAP"] = cap
    exec(compile(s2, "LightweightMemory(cap)", "exec"), g)
    return g["LightweightMemory"]


def retrieve_indexed(mem, query, top_k=3):
    """Inverted-index variant with the same scoring and tie-breaking as retrieve()."""
    qw = set(re.findall(r'\b[a-z]{3,}\b', query.lower()))
    counts = defaultdict(int)
    for w in qw:
        for idx in mem.keyword_index.get(w, ()):
            counts[idx] += 1
    best = heapq.nsmallest(top_k, counts.items(), key=lambda kv: (-kv[1], kv[0]))
    return [mem.memories[i]["text"] for i, _ in best]


def bench_memory(ns, pairs, sizes, n_queries, passes):
    rng = random.Random(SEED)
    pool = pairs[:]
    rng.shuffle(pool)
    rows = []
    for N in sizes:
        N = min(N, len(pool))
        Mem = memory_class(ns, max(N, 50))
        mem = Mem()
        entries = pool[:N]
        t0 = time.perf_counter_ns()
        for p in entries:
            mem.add_memory(p["output"])
        build_ns = time.perf_counter_ns() - t0
        assert len(mem.memories) == N, (len(mem.memories), N)
        qs = random.Random(SEED + N).sample(entries, min(n_queries, N))
        queries = [(q["instruction"] + (" " + q["input"] if q["input"] else ""), q["output"]) for q in qs]
        # correctness: identical top-3 for both variants
        same = sum(mem.retrieve(q) == retrieve_indexed(mem, q) for q, _ in queries)
        hits = sum(target in mem.retrieve(q) for q, target in queries)
        res = {"N": N, "queries": len(queries), "identical_top3": same, "recall_at_3": hits / len(queries),
               "add_mean_us": build_ns / N / 1e3,
               "mean_keywords": statistics.fmean(len(m["keywords"]) for m in mem.memories),
               "vocabulary": len(mem.keyword_index),
               "mean_postings_touched": statistics.fmean(sum(len(mem.keyword_index.get(w, ())) for w in set(re.findall(r'\b[a-z]{3,}\b', q.lower()))) for q, _ in queries)}
        for name, fn in (("scan", lambda q: mem.retrieve(q)), ("index", lambda q: retrieve_indexed(mem, q))):
            for q, _ in queries:  # warm-up pass
                fn(q)
            ts = []
            gc.collect()
            gc.disable()
            try:
                for _ in range(passes):
                    for q, _ in queries:
                        t = time.perf_counter_ns()
                        fn(q)
                        ts.append(time.perf_counter_ns() - t)
            finally:
                gc.enable()
            st = stats_ms(ts)
            res.update({f"{name}_{k}": v for k, v in st.items()})
        rows.append(res)
        log(f"memory N={N:>6}  scan median {res['scan_median_ms']:9.4f} ms  index median {res['index_median_ms']:9.4f} ms  "
            f"recall@3 {res['recall_at_3']:.3f}  identical {same}/{len(queries)}")
    # deployed configuration: unmodified class, 50-entry window, 1,000 sequential insertions
    mem = ns["LightweightMemory"]()
    add_ns = []
    for p in pool[:1000]:
        t = time.perf_counter_ns()
        mem.add_memory(p["output"])
        add_ns.append(time.perf_counter_ns() - t)
    deployed = {"window": len(mem.memories),
                "add_before_cap": stats_ms(add_ns[:50]), "add_at_cap_with_rebuild": stats_ms(add_ns[50:])}
    log(f"deployed window={deployed['window']}  add<=50: {deployed['add_before_cap']['median_ms'] * 1e3:.1f} us  "
        f"add>50 (rebuild): {deployed['add_at_cap_with_rebuild']['median_ms'] * 1e3:.1f} us")
    return rows, deployed


def run_local(args):
    out = os.path.join(args.out, "local")
    os.makedirs(out, exist_ok=True)
    if hasattr(os, "sched_setaffinity"):
        try:
            os.sched_setaffinity(0, {min(os.sched_getaffinity(0))})
        except Exception:
            pass
    env = environment(args.app)
    log("environment:", json.dumps({k: env[k] for k in ("cpu", "ram_gb", "python", "platform")}))
    _, pairs = load_alpaca(args.out)
    ns, reset, chain = load_app(args.app)
    files = make_test_files(args.out, pairs)
    parse = bench_parsing(ns, files, args.reps)
    build = bench_assembly(ns, reset, pairs, args.reps)
    sizes = [10, 50, 100, 500, 1000, 5000, 10000, 20000, len(pairs)]
    mem_rows, deployed = bench_memory(ns, pairs, sizes, args.queries, args.passes)
    summary = {"environment": env, "protocol": {"reps": args.reps, "warmup": 3, "memory_queries": args.queries,
               "memory_passes": args.passes, "seed": SEED, "alpaca_eligible_pairs": len(pairs),
               "timer": "time.perf_counter_ns, GC disabled during timed loops, process pinned to one CPU"},
               "parsing": parse, "assembly": build, "memory": mem_rows, "memory_deployed": deployed}
    json.dump(summary, open(os.path.join(out, "summary_local.json"), "w"), indent=2)
    make_local_figures(summary, out)
    log("wrote", os.path.join(out, "summary_local.json"))


# ---------------------------------------------------------------- figures (local)
INK, INK2, GRID = "#0b0b0b", "#52514e", "#e4e3df"
SERIES = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#6250d6", "#e34948"]


def _style(ax):
    ax.grid(True, which="major", color=GRID, lw=0.6)
    ax.set_axisbelow(True)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    for s in ("left", "bottom"):
        ax.spines[s].set_color(INK2)
    ax.tick_params(colors=INK2, labelsize=9)


def make_local_figures(summary, out):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9.5, "axes.labelcolor": INK, "text.color": INK})
    # Figure: component processing time
    fig, (a, b) = plt.subplots(1, 2, figsize=(7.4, 3.2), gridspec_kw={"width_ratios": [1.15, 1]})
    order = ["txt", "csv", "json", "docx", "pdf"]
    names = {"txt": "TXT", "csv": "CSV", "json": "JSON", "docx": "DOCX", "pdf": "PDF"}
    for i, t in enumerate(order):
        r = [x for x in summary["parsing"] if x["type"] == t]
        xs = [x["bytes"] / 1024 for x in r]
        ys = [x["median_ms"] for x in r]
        a.plot(xs, ys, "-o", color=SERIES[i], lw=2, ms=5, label=names[t])
        a.annotate(names[t], (xs[-1], ys[-1]), xytext=(5, 0), textcoords="offset points", va="center", fontsize=8.5, color=INK)
    a.set_xscale("log"); a.set_yscale("log")
    a.set_xlabel("File size (KB, log scale)"); a.set_ylabel("Median parsing time (ms, log scale)")
    a.set_xlim(right=max(x["bytes"] for x in summary["parsing"]) / 1024 * 4)
    _style(a)
    a.set_title("(a) File parsing", fontsize=10, y=-0.36)
    lab = {"PowerPoint (.pptx)": "PPTX", "Word (.docx)": "DOCX"}
    rows = summary["assembly"]
    short = {"phone sales": "sales", "student data": "students", "project tasks": "generic"}
    labels = [lab[r["artifact"]] if r["artifact"] in lab else "XLSX\n" + short.get(r["artifact"].split(": ")[1], r["artifact"].split(": ")[1]) for r in rows]
    med = [r["median_ms"] for r in rows]
    err = [[r["median_ms"] - r["q1_ms"] for r in rows], [r["q3_ms"] - r["median_ms"] for r in rows]]
    b.bar(range(len(rows)), med, color=SERIES[0], width=0.6, yerr=err, ecolor=INK2, capsize=3)
    for i, v in enumerate(med):
        b.text(i, rows[i]["q3_ms"] + max(med) * 0.03, f"{v:.1f}", ha="center", va="bottom", fontsize=8, color=INK)
    b.set_xticks(range(len(rows)), labels, fontsize=7.6)
    b.set_ylabel("Median build time (ms)")
    b.set_ylim(0, max(med) * 1.25)
    _style(b)
    b.grid(False, axis="x")
    b.set_title("(b) Document assembly", fontsize=10, y=-0.36)
    fig.tight_layout(w_pad=2.5)
    fig.savefig(os.path.join(out, "fig_components.png"), dpi=300, bbox_inches="tight")
    plt.close(fig)
    # Figure: memory scaling
    m = summary["memory"]
    N = [r["N"] for r in m]
    fig, (a, b) = plt.subplots(1, 2, figsize=(7.4, 3.2))
    for i, (k, name) in enumerate((("scan", "Linear scan (deployed)"), ("index", "Inverted index (proposed)"))):
        med = [r[f"{k}_median_ms"] for r in m]
        p95 = [r[f"{k}_p95_ms"] for r in m]
        a.fill_between(N, med, p95, color=SERIES[i], alpha=0.15, lw=0)
        a.plot(N, med, "-o", color=SERIES[i], lw=2, ms=5, label=name)
    a.axvline(50, color=INK2, lw=1, ls="--")
    a.text(50 * 1.2, max(r["scan_p95_ms"] for r in m) * 1.6, "deployed window (N = 50)", fontsize=8, color=INK2, va="top")
    a.set_xscale("log"); a.set_yscale("log")
    a.set_xlabel("Stored entries N (log scale)"); a.set_ylabel("Retrieval latency per query (ms)")
    a.legend(frameon=False, fontsize=8, loc="lower right")
    _style(a)
    a.set_title("(a) Retrieval latency (median, band to p95)", fontsize=10, y=-0.36)
    b.plot(N, [r["recall_at_3"] * 100 for r in m], "-o", color=SERIES[0], lw=2, ms=5)
    b.axvline(50, color=INK2, lw=1, ls="--")
    b.set_xscale("log")
    b.set_ylim(0, 100)
    b.set_xlabel("Stored entries N (log scale)"); b.set_ylabel("Recall@3 (%)")
    _style(b)
    b.set_title("(b) Retrieval quality", fontsize=10, y=-0.36)
    fig.tight_layout(w_pad=2.5)
    fig.savefig(os.path.join(out, "fig_memory.png"), dpi=300, bbox_inches="tight")
    plt.close(fig)


# ---------------------------------------------------------------- ONLINE benchmark
WEB_QUERIES = [
    "What is the population of Seoul?", "Who is the current president of South Korea?",
    "What is the tallest building in the world?", "Tell me about Kyungdong University.",
    "What is the weather in Seoul?", "Who is the CEO of Samsung Electronics?",
    "What is the latest version of Python?", "Tell me about the Nobel Prize in Physics.",
    "What is the largest lake in Africa?", "Tell me about the Hubble Space Telescope."]
FILE_QUESTIONS = ["Summarize the main points of the uploaded file."]
EXCEL_PROMPTS = ["generate an excel about phone sales", "make an excel about student data",
                 "create an excel about product inventory", "generate an excel about monthly expenses",
                 "make an excel about library loans", "create an excel about sales targets",
                 "generate an excel about student attendance", "make an excel about warehouse products",
                 "create an excel about phone sales by region", "generate an excel about project tasks"]
PPT_PROMPTS = ["make a ppt about artificial intelligence", "create a powerpoint about renewable energy",
               "make a ppt about cloud computing", "create a ppt about cybersecurity basics",
               "make a powerpoint about smart cities"]
WORD_PROMPTS = ["create a word document about machine learning", "make a word document about climate change",
                "create a word document about digital health", "make a word document about internet of things",
                "create a word document about online education"]
IMAGE_PROMPTS = ["generate image of a red bicycle in the rain", "generate image of a mountain lake at sunrise",
                 "generate image of a robot reading a book", "generate image of a lighthouse at night",
                 "generate image of a bowl of fresh fruit", "generate image of a city street in winter",
                 "generate image of a cat sleeping on a sofa", "generate image of a sailing boat on the sea",
                 "generate image of a forest path in autumn", "generate image of a desert under stars"]
ONLINE_FILES = ["report_10pages.pdf", "document_50para.docx", "text_100KB.txt", "table_1000rows.csv", "records_1000.json"]


def build_plan(pairs, files):
    rng = random.Random(SEED)
    pool = [p for p in pairs if not p["input"] and len(p["instruction"]) < 120]
    llm_q = [p["instruction"] for p in rng.sample(pool, 10)]
    plan = [("llm", i, q) for i, q in enumerate(llm_q)]
    plan += [("web", i, q) for i, q in enumerate(WEB_QUERIES)]
    fmap = {f["name"]: f for f in files}
    plan += [("file", i, (fmap[n], FILE_QUESTIONS[0])) for i, n in enumerate(ONLINE_FILES)]
    plan += [("excel", i, q) for i, q in enumerate(EXCEL_PROMPTS)]
    plan += [("ppt", i, q) for i, q in enumerate(PPT_PROMPTS)]
    plan += [("word", i, q) for i, q in enumerate(WORD_PROMPTS)]
    plan += [("image", i, q) for i, q in enumerate(IMAGE_PROMPTS)]
    return plan


def check_routes(ns, reset, plan):
    for op, i, q in plan:
        reset()
        if op == "web":
            r = ns["route"](q)
            assert r == "search", (q, r)
        if op == "file":
            assert ns["route"](q[1]) == "file_task", q[1]
        if op == "image":
            assert ns["route"](q) == "generate_image", q


def classify(e):
    code = getattr(e, "status_code", None)
    msg = str(e).lower()
    name = type(e).__name__
    if "empty completion" in msg:
        return "empty_completion"
    if code == 429 or "rate limit" in msg:
        return "rate_limit_daily" if ("per day" in msg or "(tpd)" in msg or "(rpd)" in msg) else "rate_limit"
    if "decommission" in msg or "does not exist" in msg or code == 404:
        return "model_unavailable"
    if code and code >= 500:
        return "server_error"
    if "timeout" in name.lower() or "timed out" in msg:
        return "timeout"
    if "connection" in name.lower():
        return "connection"
    return f"{name}:{code}" if code else name


class Recorder:
    def __init__(self):
        self.attempts, self.requests, self.searches = [], [], []
        self.op_id, self.req_id, self.req_counter = None, None, 0


class InstrumentedClient:
    def __init__(self, inner, rec):
        self._inner, self._rec = inner, rec
        self.chat = types.SimpleNamespace(completions=types.SimpleNamespace(create=self._create))

    def _create(self, **kw):
        rec = self._rec
        row = {"op_id": rec.op_id, "req_id": rec.req_id, "model": kw.get("model"), "wall_start": time.time()}
        t0 = time.perf_counter()
        try:
            r = self._inner.chat.completions.create(**kw)
            content = r.choices[0].message.content or ""
            u = getattr(r, "usage", None)
            row.update(ok_api=1, empty=int(not content.strip()), error_type="empty_completion" if not content.strip() else "",
                       error="", prompt_tokens=getattr(u, "prompt_tokens", None), completion_tokens=getattr(u, "completion_tokens", None))
            return r
        except Exception as e:
            row.update(ok_api=0, empty=0, error_type=classify(e), error=str(e)[:200].replace("\n", " "),
                       prompt_tokens=None, completion_tokens=None)
            raise
        finally:
            row["latency_s"] = round(time.perf_counter() - t0, 4)
            rec.attempts.append(row)


def instrument(ns, rec, chain):
    orig_llm, orig_search = ns["llm_with_fallback"], ns["internet_search"]

    def llm_with_fallback(messages, max_retries=2):
        rec.req_counter += 1
        rec.req_id = rec.req_counter
        n0 = len(rec.attempts)
        t0 = time.perf_counter()
        out = orig_llm(messages, max_retries=max_retries)
        dt = time.perf_counter() - t0
        atts = rec.attempts[n0:]
        served = next((a["model"] for a in reversed(atts) if a["ok_api"] and not a["empty"]), "")
        if out == UNAVAILABLE:
            served = ""
        pos = chain.index(served) + 1 if served in chain else 0
        rec.requests.append({"op_id": rec.op_id, "req_id": rec.req_id, "latency_s": round(dt, 4), "attempts": len(atts),
                             "served_model": served, "chain_position": pos,
                             "outcome": "primary" if pos == 1 else ("fallback" if pos > 1 else "failed"),
                             "first_attempt_success": int(len(atts) == 1 and pos == 1),
                             "prompt_tokens": sum(a["prompt_tokens"] or 0 for a in atts),
                             "completion_tokens": sum(a["completion_tokens"] or 0 for a in atts)})
        rec.req_id = None
        return out

    def internet_search(query):
        t0 = time.perf_counter()
        r = orig_search(query)
        rec.searches.append({"op_id": rec.op_id, "hit": int(bool(r)), "latency_s": round(time.perf_counter() - t0, 4)})
        return r

    ns["llm_with_fallback"] = llm_with_fallback
    ns["internet_search"] = internet_search


# ----- simulated services for --dry-run
class _Msg:
    def __init__(self, c):
        self.message = types.SimpleNamespace(content=c)


class FakeGroq:
    class RateLimit(Exception):
        status_code = 429

    def __init__(self, seed=7):
        self.rng = random.Random(seed)
        self.chat = types.SimpleNamespace(completions=types.SimpleNamespace(create=self.create))

    def create(self, **kw):
        time.sleep(min(self.rng.lognormvariate(-2.5, 0.4), 0.3))
        u = self.rng.random()
        if kw["model"].endswith("120b") and u < 0.08:
            raise FakeGroq.RateLimit("Rate limit reached for model (TPM): simulated")
        if u > 0.99:
            return types.SimpleNamespace(choices=[_Msg("")], usage=None)
        text = PPT_CONTENT if "PowerPoint" in kw["messages"][-1]["content"] else "Simulated answer with enough text to be stored in memory for testing."
        return types.SimpleNamespace(choices=[_Msg(text)], usage=types.SimpleNamespace(prompt_tokens=900, completion_tokens=300))

    class models:
        @staticmethod
        def list():
            return types.SimpleNamespace(data=[types.SimpleNamespace(id=m) for m in ("openai/gpt-oss-120b", "openai/gpt-oss-20b")])


class _Resp:
    def __init__(self, status=200, text="", content=b"", ctype="text/html"):
        self.status_code, self.text, self.content, self.headers = status, text, content, {"Content-Type": ctype}

    def raise_for_status(self):
        pass


def fake_requests(real):
    html = ''.join(f'<a rel="nofollow" class="result__a" href="https://example.org/{i}">Result {i}</a>'
                   f'<a class="result__snippet" href="https://example.org/{i}">Snippet text number {i}.</a>' for i in range(3))

    def post(url, **k):
        time.sleep(0.05)
        return _Resp(200, html)

    def get(url, **k):
        time.sleep(0.05)
        if "pollinations" in url:
            return _Resp(200, "", b"\x89PNG" + b"0" * 2048, "image/jpeg")
        return _Resp(200, "<html><body>" + "content " * 100 + "</body></html>", b"x")
    return types.SimpleNamespace(post=post, get=get, utils=real.utils, exceptions=real.exceptions)


def fetch_image(http, url, timeout):
    t0 = time.perf_counter()
    try:
        r = http.get(url, timeout=timeout)
        ok = r.status_code == 200 and r.headers.get("Content-Type", "").startswith("image") and len(r.content) > 0
        return ok, len(r.content), time.perf_counter() - t0, "" if ok else f"status {r.status_code} {r.headers.get('Content-Type', '')}"
    except Exception as e:
        return False, 0, time.perf_counter() - t0, type(e).__name__


class QuotaExhausted(Exception):
    pass


def run_op(ns, rec, op, payload, http, img_timeout):
    S = ns["st"].session_state
    res = {"parse_s": "", "fetch_s": "", "out_bytes": "", "detail": ""}
    t0 = time.perf_counter()
    if op == "llm":
        out = ns["reason"](payload, "")
        ok = bool(out.strip()) and out != UNAVAILABLE
    elif op == "web":
        out = ns["run_agent"](payload)
        ok = bool(out.strip()) and out != UNAVAILABLE
    elif op == "file":
        f, q = payload
        tp = time.perf_counter()
        content = ns["process_uploaded_file"](Upload(open(f["path"], "rb").read(), f["name"], MIME[f["type"]]))
        res["parse_s"] = round(time.perf_counter() - tp, 4)
        S.uploaded_files = {f["name"]: content}  # same assembly as the sidebar uploader
        S.file_context = "\n".join(f"\n{'=' * 50}\n📄 {n}\n{'=' * 50}\n{c}\n" for n, c in S.uploaded_files.items())
        out = ns["run_agent"](q)
        ok = bool(out.strip()) and out != UNAVAILABLE
    elif op in ("excel", "ppt", "word"):
        out = ns["run_agent"](payload)
        b = S.get({"excel": "excel_data", "ppt": "ppt_data", "word": "word_data"}[op])
        ok = b is not None
        res["out_bytes"] = len(b.getvalue()) if b is not None else 0
    elif op == "image":
        out = ns["run_agent"](payload)
        m = re.search(r"!\[[^\]]*\]\((https?://[^)\s]+)\)", out)
        handler = time.perf_counter() - t0
        if m:
            ok, nbytes, fetch, err = fetch_image(http, m.group(1), img_timeout)
            res.update(fetch_s=round(fetch, 4), out_bytes=nbytes, detail=err)
        else:
            ok, res["detail"] = False, "no image URL"
        res["handler_s"] = round(handler, 4)
    total = time.perf_counter() - t0
    reqs = [r for r in rec.requests if r["op_id"] == rec.op_id]
    atts = [a for a in rec.attempts if a["op_id"] == rec.op_id]
    if any(a["error_type"] == "rate_limit_daily" for a in atts):
        raise QuotaExhausted()
    if reqs and any(r["outcome"] == "failed" for r in reqs):
        ok = False
    srch = [s for s in rec.searches if s["op_id"] == rec.op_id]
    res.update(total_s=round(total, 4), llm_s=round(sum(r["latency_s"] for r in reqs), 4), n_llm_requests=len(reqs),
               search_hit=srch[-1]["hit"] if srch else "", success=int(ok))
    res.setdefault("handler_s", round(total - (res["parse_s"] or 0), 4))
    return res


def append_csv(path, rows):
    if not rows:
        return
    new = not os.path.exists(path)
    with open(path, "a", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        if new:
            w.writeheader()
        w.writerows(rows)


def read_csv(path):
    if not os.path.exists(path):
        return []
    return list(csv.DictReader(open(path, encoding="utf-8")))


def preflight(ns, chain, client, http, dry):
    log(f"fallback chain read from app.py: {chain}")
    if not chain:
        sys.exit("models_to_try not found in llm_with_fallback()")
    if not dry:
        avail = {m.id for m in client.models.list().data}
        missing = [m for m in chain if m not in avail]
        if missing:
            sys.exit(f"ABORT: chain models not served to this API key: {missing}. Run make_ijasc_build.py or edit models_to_try.")
        for m in chain:
            t0 = time.perf_counter()
            r = client.chat.completions.create(model=m, messages=[{"role": "user", "content": "Reply with the single word OK."}],
                                               max_tokens=64, temperature=0, timeout=30,
                                               extra_body={"reasoning_effort": "low"} if m.startswith("openai/gpt-oss") else None)
            log(f"  {m}: {time.perf_counter() - t0:.2f}s reply={(r.choices[0].message.content or '').strip()[:20]!r}")
    t0 = time.perf_counter()
    s = ns["internet_search"]("What is the population of Seoul?")
    log(f"  DuckDuckGo HTML search: {'OK' if s else 'NO RESULTS (search ops will fall back to date context)'} in {time.perf_counter() - t0:.2f}s")
    ok, n, dt, err = fetch_image(http, "https://image.pollinations.ai/prompt/test%20pattern?width=64&height=64&nologo=true", 60)
    log(f"  Pollinations image endpoint: {'OK' if ok else 'FAILED ' + err} ({n} bytes, {dt:.1f}s)")


def summarize_online(out, chain, env, args):
    ops, reqs, atts = read_csv(os.path.join(out, "ops.csv")), read_csv(os.path.join(out, "requests.csv")), read_csv(os.path.join(out, "attempts.csv"))
    f = lambda v: float(v) if v not in ("", None) else None
    by = defaultdict(list)
    for o in ops:
        by[o["op"]].append(o)
    names = {"llm": "LLM response (no retrieval)", "web": "Web search + LLM response", "file": "File analysis (parse + LLM)",
             "excel": "Excel generation", "ppt": "PowerPoint generation", "word": "Word generation", "image": "Image generation (URL + download)"}
    table = []
    for k in ("llm", "web", "file", "excel", "ppt", "word", "image"):
        rows = by.get(k, [])
        if not rows:
            continue
        tot = [f(r["total_s"]) for r in rows]
        okr = [f(r["total_s"]) for r in rows if r["success"] == "1"]
        llm = [f(r["llm_s"]) for r in rows if f(r["n_llm_requests"])]
        e = {"op": k, "label": names[k], "n": len(rows), "success": sum(r["success"] == "1" for r in rows),
             "median_s": statistics.median(tot), "mean_s": statistics.fmean(tot), "q1_s": pctl(tot, .25), "q3_s": pctl(tot, .75),
             "p95_s": pctl(tot, .95), "median_llm_s": statistics.median(llm) if llm else None,
             "median_success_only_s": statistics.median(okr) if okr else None}
        if k == "web":
            hits = [r["search_hit"] for r in rows if r["search_hit"] != ""]
            e["search_hit_rate"] = sum(h == "1" for h in hits) / len(hits) if hits else None
        if k == "file":
            e["median_parse_s"] = statistics.median(f(r["parse_s"]) for r in rows)
        if k == "image":
            fx = [f(r["fetch_s"]) for r in rows if r["fetch_s"] != ""]
            e["median_fetch_s"] = statistics.median(fx) if fx else None
        table.append(e)
    n = len(reqs)
    cnt = lambda o: sum(r["outcome"] == o for r in reqs)
    rel = {"definitions": {
        "request": "one call of llm_with_fallback() (one user-visible LLM answer)",
        "attempt": "one chat.completions.create() call to one model (the SDK's own internal retries are inside one attempt)",
        "primary_served": "request answered by chain position 1 (on its first or second attempt)",
        "fallback_served": "request answered by chain position >= 2",
        "failed": "no model in the chain returned a non-empty completion; the user sees the unavailable message",
        "availability": "(primary_served + fallback_served) / requests",
        "fallback_rate": "fallback_served / requests"},
        "chain": chain, "requests": n, "primary_served": cnt("primary"), "fallback_served": cnt("fallback"), "failed": cnt("failed"),
        "first_attempt_success": sum(r["first_attempt_success"] == "1" for r in reqs)}
    if n:
        rel.update(availability=(cnt("primary") + cnt("fallback")) / n, fallback_rate=cnt("fallback") / n,
                   primary_rate=cnt("primary") / n, failure_rate=cnt("failed") / n,
                   first_attempt_success_rate=rel["first_attempt_success"] / n,
                   attempts_total=len(atts), attempts_per_request=len(atts) / n,
                   tokens_total=sum(int(r["prompt_tokens"]) + int(r["completion_tokens"]) for r in reqs))
    per_model = []
    for m in chain:
        a = [x for x in atts if x["model"] == m]
        succ = sum(x["ok_api"] == "1" and x["empty"] == "0" for x in a)
        lat = [f(x["latency_s"]) for x in a if x["ok_api"] == "1" and x["empty"] == "0"]
        per_model.append({"model": m, "attempts": len(a), "successful_attempts": succ,
                          "attempt_success_rate": succ / len(a) if a else None,
                          "requests_served": sum(r["served_model"] == m for r in reqs),
                          "median_success_latency_s": statistics.median(lat) if lat else None})
    rel["per_model"] = per_model
    causes = defaultdict(int)
    for x in atts:
        if x["error_type"]:
            causes[x["error_type"]] += 1
    rel["failed_attempt_causes"] = dict(causes)
    summary = {"environment": env, "protocol": {"reps": args.reps, "gap_s": args.gap, "seed": SEED, "image_timeout_s": args.image_timeout,
               "fresh_session_per_operation": True, "order": "operations shuffled per repetition with seed 2026+rep",
               "dry_run": args.dry_run}, "operations": table, "reliability": rel}
    json.dump(summary, open(os.path.join(out, "summary_online.json"), "w"), indent=2)
    L = ["=" * 78, "MozeAI ONLINE BENCHMARK SUMMARY" + ("  [DRY RUN - SIMULATED SERVICES]" if args.dry_run else ""), "=" * 78,
         f"env: {env['cpu']} | {env['ram_gb']} GB | Python {env['python']} | {env['platform']}",
         f"app sha256 {env['app_sha256'][:12]} | git {env.get('git_describe', '?')} | groq {env['packages'].get('groq')}",
         f"reps={args.reps} gap={args.gap}s fresh session per op, shuffled order (seed {SEED})", "",
         f"{'operation':<36}{'n':>4}{'ok':>5}{'median':>9}{'IQR':>16}{'p95':>8}{'LLM med':>9}"]
    for e in table:
        L.append(f"{e['label']:<36}{e['n']:>4}{e['success']:>5}{e['median_s']:>8.2f}s{e['q1_s']:>8.2f}-{e['q3_s']:<6.2f}{e['p95_s']:>7.2f}s"
                 f"{(e['median_llm_s'] or 0):>8.2f}s")
        extra = {k: round(v, 3) for k, v in e.items() if k in ("search_hit_rate", "median_parse_s", "median_fetch_s") and v is not None}
        if extra:
            L.append(f"{'':<36}{extra}")
    L += ["", f"LLM requests N={n}: primary={rel['primary_served']} fallback={rel['fallback_served']} failed={rel['failed']}"]
    if n:
        L.append(f"availability={rel['availability']:.4f} fallback_rate={rel['fallback_rate']:.4f} first_attempt_success={rel['first_attempt_success_rate']:.4f} "
                 f"attempts/request={rel['attempts_per_request']:.3f} tokens={rel['tokens_total']}")
    for pm in per_model:
        L.append(f"  {pm['model']:<24} attempts={pm['attempts']:<4} ok={pm['successful_attempts']:<4} served={pm['requests_served']:<4} "
                 f"success_rate={pm['attempt_success_rate'] if pm['attempt_success_rate'] is None else round(pm['attempt_success_rate'], 4)} "
                 f"median_latency={pm['median_success_latency_s']}")
    L.append(f"failed-attempt causes: {dict(causes) or 'none'}")
    L.append("=" * 78)
    txt = "\n".join(L)
    open(os.path.join(out, "summary_online.txt"), "w").write(txt + "\n")
    print(txt, flush=True)


def run_online(args):
    out = os.path.join(args.out, "online_dryrun" if args.dry_run else "online")
    os.makedirs(out, exist_ok=True)
    env = environment(args.app)
    _, pairs = load_alpaca(args.out)
    files = make_test_files(args.out, pairs)
    rec = Recorder()
    if args.dry_run:
        inner = FakeGroq()
    else:
        key = os.environ.get("GROQ_API_KEY")
        if not key:
            sys.exit("set GROQ_API_KEY first (export GROQ_API_KEY=...)")
        import groq
        inner = groq.Groq(api_key=key)  # same construction as app.py (SDK default retry policy)
    ns, reset, chain = load_app(args.app, InstrumentedClient(inner, rec))
    import requests as real_requests
    http = fake_requests(real_requests) if args.dry_run else real_requests
    if args.dry_run:
        ns["requests"] = http
    instrument(ns, rec, chain)
    plan = build_plan(pairs, files)
    check_routes(ns, reset, plan)
    preflight(ns, chain, inner, http, args.dry_run)
    rec.attempts.clear(); rec.requests.clear(); rec.searches.clear()
    n_llm = sum(op in ("llm", "web", "file", "ppt", "word") for op, _, _ in plan) * args.reps
    log(f"plan: {len(plan)} operations x {args.reps} reps = {len(plan) * args.reps} runs, ~{n_llm} LLM requests "
        f"(~{n_llm * 1.6:.0f}K tokens), gap {args.gap}s -> roughly {len(plan) * args.reps * (args.gap + 4) / 60:.0f} min")
    if n_llm * 1.6 > 180 and not args.dry_run:
        log("NOTE: close to the free-tier daily limit (200K tokens per model). If the run stops with a daily-quota "
            "message, rerun the same command after the reset; completed runs are kept and skipped.")
    if args.preflight_only:
        return
    done = {(r["op"], r["item"], r["rep"]) for r in read_csv(os.path.join(out, "ops.csv"))}
    if done:
        log(f"resuming: {len(done)} runs already recorded in {out}/ops.csv")
    k, total = 0, len(plan) * args.reps
    for rep in range(args.reps):
        order = plan[:]
        random.Random(SEED + rep).shuffle(order)
        for op, item, payload in order:
            k += 1
            if (op, str(item), str(rep)) in done:
                continue
            reset()
            rec.op_id = f"{rep}:{op}:{item}"
            n_a, n_r, n_s = len(rec.attempts), len(rec.requests), len(rec.searches)
            try:
                res = run_op(ns, rec, op, payload, http, args.image_timeout)
            except QuotaExhausted:
                del rec.attempts[n_a:], rec.requests[n_r:], rec.searches[n_s:]
                log("STOP: Groq daily quota exhausted. Results so far are saved; rerun the same command later to resume.")
                summarize_online(out, chain, env, args)
                sys.exit(3)
            q = payload[0]["name"] + " | " + payload[1] if op == "file" else payload
            row = {"op_id": rec.op_id, "rep": rep, "op": op, "item": item, "query": q, **{x: res.get(x, "") for x in
                   ("total_s", "handler_s", "parse_s", "fetch_s", "llm_s", "n_llm_requests", "search_hit", "out_bytes", "success", "detail")}}
            append_csv(os.path.join(out, "ops.csv"), [row])
            append_csv(os.path.join(out, "requests.csv"), rec.requests[n_r:])
            append_csv(os.path.join(out, "attempts.csv"), rec.attempts[n_a:])
            served = ",".join(r["served_model"].split("/")[-1] or "FAILED" for r in rec.requests[n_r:]) or "-"
            log(f"[{k:>3}/{total}] rep{rep} {op:<5} #{item:<2} {res['total_s']:7.2f}s llm {res['llm_s']:6.2f}s served={served:<14} ok={res['success']} {res['detail']}")
            time.sleep(args.gap)
    summarize_online(out, chain, env, args)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("mode", choices=["local", "online", "figures"])
    ap.add_argument("--app", default="emotion-ai-demo/app.py")
    ap.add_argument("--out", default="bench_results")
    ap.add_argument("--reps", type=int, default=None, help="local: timed repetitions per case (30); online: repetitions of the plan (3)")
    ap.add_argument("--queries", type=int, default=200, help="local memory benchmark: queries per collection size")
    ap.add_argument("--passes", type=int, default=5, help="local memory benchmark: timed passes over the query set")
    ap.add_argument("--gap", type=float, default=15.0, help="online: pause between operations, seconds (free tier: 8K tokens/min)")
    ap.add_argument("--image-timeout", type=float, default=120.0)
    ap.add_argument("--dry-run", action="store_true", help="online: simulated Groq/DuckDuckGo/Pollinations (self-test)")
    ap.add_argument("--preflight-only", action="store_true")
    args = ap.parse_args()
    if args.reps is None:
        args.reps = 30 if args.mode == "local" else 3
    if args.dry_run:
        args.gap = min(args.gap, 0.0)
    if args.mode == "figures":  # redraw local figures from an existing summary_local.json
        out = os.path.join(args.out, "local")
        make_local_figures(json.load(open(os.path.join(out, "summary_local.json"))), out)
        return
    (run_local if args.mode == "local" else run_online)(args)


if __name__ == "__main__":
    main()
